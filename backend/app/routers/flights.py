from fastapi import APIRouter, Depends, File, HTTPException, UploadFile

from app.database import get_connection
from app.deps import get_current_user
from app.schemas import CostResultOut, FlightCreate, FlightOut, FlightUpdate
from app.services.cost import calculate_for_flight, get_cost_result

router = APIRouter(prefix="/flights", tags=["flights"])

_COLUMNS = (
    "flight_number, departure_airport, arrival_airport, distance_km, duration_hours, "
    "aircraft_id, passengers, fuel_price_per_liter, ticket_price_avg, catering_cost_per_pax, "
    "handling_cost, taxes_airport, flight_date, status"
)
_PLACEHOLDERS = "%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s"


def _to_dict(r) -> dict:
    d = dict(r)
    flight_date = d["flight_date"]
    if flight_date is not None and not isinstance(flight_date, str):
        # PostgreSQL renvoie un objet date, on le convertit en ISO pour l'API.
        flight_date = flight_date.isoformat()
    return {
        "id": d["id"],
        "flight_number": d["flight_number"],
        "departure_airport": d["departure_airport"],
        "arrival_airport": d["arrival_airport"],
        "distance_km": d["distance_km"],
        "duration_hours": d["duration_hours"],
        "aircraft_id": d["aircraft_id"],
        "passengers": d["passengers"],
        "fuel_price_per_liter": d["fuel_price_per_liter"],
        "ticket_price_avg": d["ticket_price_avg"],
        "catering_cost_per_pax": d["catering_cost_per_pax"],
        "handling_cost": d["handling_cost"],
        "taxes_airport": d["taxes_airport"],
        "flight_date": flight_date,
        "status": d["status"],
        "aircraft_model": d.get("aircraft_model"),
        "departure_airport_name": d.get("departure_airport_name"),
        "departure_airport_city": d.get("departure_airport_city"),
        "arrival_airport_name": d.get("arrival_airport_name"),
        "arrival_airport_city": d.get("arrival_airport_city"),
        "total_cost": d.get("total_cost"),
        "total_revenue": d.get("total_revenue"),
        "profit_margin": d.get("profit_margin"),
        "is_profitable": bool(d["is_profitable"]) if d.get("is_profitable") is not None else None,
    }


def _body_values(body) -> tuple:
    return (
        body.flight_number, body.departure_airport, body.arrival_airport,
        body.distance_km, body.duration_hours, body.aircraft_id, body.passengers,
        body.fuel_price_per_liter, body.ticket_price_avg, body.catering_cost_per_pax,
        body.handling_cost, body.taxes_airport, body.flight_date, body.status,
    )


@router.get("", response_model=list[FlightOut])
def list_flights(_=Depends(get_current_user)):
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT f.*, a.model AS aircraft_model, "
            "dep.name AS departure_airport_name, dep.city AS departure_airport_city, "
            "arr.name AS arrival_airport_name, arr.city AS arrival_airport_city, "
            "cr.total_cost, cr.total_revenue, cr.profit_margin, cr.is_profitable "
            "FROM flights f "
            "LEFT JOIN aircraft a ON f.aircraft_id = a.id "
            "LEFT JOIN airports dep ON f.departure_airport = dep.code "
            "LEFT JOIN airports arr ON f.arrival_airport = arr.code "
            "LEFT JOIN cost_results cr ON cr.flight_id = f.id "
            "ORDER BY f.created_at DESC, f.id ASC"
        ).fetchall()
        return [_to_dict(r) for r in rows]
    finally:
        conn.close()


def _flight_row(conn, flight_id: int):
    """Renvoie un vol enrichi (avion + aéroports de départ/arrivée)."""
    return conn.execute(
        "SELECT f.*, a.model AS aircraft_model, "
        "dep.name AS departure_airport_name, dep.city AS departure_airport_city, "
        "arr.name AS arrival_airport_name, arr.city AS arrival_airport_city "
        "FROM flights f "
        "LEFT JOIN aircraft a ON f.aircraft_id = a.id "
        "LEFT JOIN airports dep ON f.departure_airport = dep.code "
        "LEFT JOIN airports arr ON f.arrival_airport = arr.code "
        "WHERE f.id=%s", (flight_id,)
    ).fetchone()


def _validate_airport_ref(conn, code: str) -> None:
    """Vérifie qu'un code IATA existe, sinon lève une erreur 400 explicite."""
    row = conn.execute("SELECT code FROM airports WHERE code=%s", (code,)).fetchone()
    if row is None:
        raise HTTPException(400, f"Code IATA '{code}' introuvable. Ajoutez-le dans la gestion des aeroports.")


@router.get("/{flight_id}", response_model=FlightOut)
def get_flight(flight_id: int, _=Depends(get_current_user)):
    conn = get_connection()
    try:
        row = _flight_row(conn, flight_id)
    finally:
        conn.close()
    if row is None:
        raise HTTPException(404, "Vol introuvable")
    return _to_dict(row)


@router.post("", response_model=FlightOut)
def create_flight(body: FlightCreate, _=Depends(get_current_user)):
    conn = get_connection()
    try:
        _validate_airport_ref(conn, body.departure_airport)
        _validate_airport_ref(conn, body.arrival_airport)
        cur = conn.execute(
            f"INSERT INTO flights ({_COLUMNS}) VALUES ({_PLACEHOLDERS}) RETURNING id",
            _body_values(body),
        )
        new_id = cur.fetchone()["id"]
        conn.commit()
        row = _flight_row(conn, new_id)
    finally:
        conn.close()
    return _to_dict(row)


@router.put("/{flight_id}", response_model=FlightOut)
def update_flight(flight_id: int, body: FlightUpdate, _=Depends(get_current_user)):
    conn = get_connection()
    try:
        exists = conn.execute("SELECT id FROM flights WHERE id=%s", (flight_id,)).fetchone()
        if exists is None:
            raise HTTPException(404, "Vol introuvable")
        _validate_airport_ref(conn, body.departure_airport)
        _validate_airport_ref(conn, body.arrival_airport)
        conn.execute(
            "UPDATE flights SET flight_number=%s, departure_airport=%s, arrival_airport=%s, "
            "distance_km=%s, duration_hours=%s, aircraft_id=%s, passengers=%s, "
            "fuel_price_per_liter=%s, ticket_price_avg=%s, catering_cost_per_pax=%s, "
            "handling_cost=%s, taxes_airport=%s, flight_date=%s, status=%s WHERE id=%s",
            (*_body_values(body), flight_id),
        )
        conn.commit()
        row = _flight_row(conn, flight_id)
        return _to_dict(row)
    finally:
        conn.close()


@router.delete("/{flight_id}")
def delete_flight(flight_id: int, _=Depends(get_current_user)):
    conn = get_connection()
    try:
        conn.execute("DELETE FROM cost_results WHERE flight_id=%s", (flight_id,))
        conn.execute("DELETE FROM simulations WHERE flight_id=%s", (flight_id,))
        conn.execute("DELETE FROM flights WHERE id=%s", (flight_id,))
        conn.commit()
    finally:
        conn.close()
    return {"ok": True}


@router.post("/{flight_id}/calculate", response_model=CostResultOut)
def calculate_flight(flight_id: int, _=Depends(get_current_user)):
    try:
        return calculate_for_flight(flight_id)
    except ValueError as e:
        raise HTTPException(400, str(e))


@router.get("/{flight_id}/cost", response_model=CostResultOut)
def flight_cost(flight_id: int, _=Depends(get_current_user)):
    result = get_cost_result(flight_id)
    if result is None:
        raise HTTPException(404, "Cout non calcule pour ce vol")
    return result


def _read_excel_bytes(content: bytes) -> dict:
    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb["Vol"] if "Vol" in wb.sheetnames else wb.active
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        champ = row[0] if len(row) > 0 else None
        valeur = row[1] if len(row) > 1 else None
        if not champ:
            continue
        champ_str = str(champ).strip()
        if champ_str.startswith("---") or champ_str == "":
            continue
        data[champ_str] = valeur
    return data


def _validate_excel(data: dict) -> list:
    errors = []
    obligatoires = [
        "Numero de vol", "Aeroport depart", "Aeroport arrivee", "Distance (km)",
        "Duree (heures)", "Nombre de passagers", "Prix carburant (MAD/L)",
        "Prix billet moyen (MAD)", "ID Avion",
    ]
    for champ in obligatoires:
        if champ not in data or data[champ] is None or str(data[champ]).strip() == "":
            errors.append("Champ manquant : " + champ)

    champs_numeriques = [
        "Distance (km)",
        "Duree (heures)",
        "Nombre de passagers",
        "Prix carburant (MAD/L)",
        "Prix billet moyen (MAD)",
    ]
    for champ in champs_numeriques:
        if champ in data and data[champ] is not None:
            try:
                val = float(str(data[champ]).replace(",", "."))
                if val <= 0:
                    errors.append(champ + " doit etre superieur a 0")
            except ValueError:
                errors.append(champ + " doit etre un nombre")

    if "ID Avion" in data and data["ID Avion"] is not None:
        try:
            aid = int(float(str(data["ID Avion"])))
            conn = get_connection()
            try:
                exists = conn.execute("SELECT id FROM aircraft WHERE id=%s", (aid,)).fetchone()
            finally:
                conn.close()
            if exists is None:
                errors.append(f"ID Avion {aid} introuvable dans la base de donnees")
        except ValueError:
            errors.append("ID Avion doit etre un nombre entier")

    # Validation des codes IATA (depart & arrivee) : ils doivent exister dans
    # la table airports pour garantir la coherence des donnees.
    for champ in ("Aeroport depart", "Aeroport arrivee"):
        if champ in data and data[champ] is not None and str(data[champ]).strip() != "":
            code = str(data[champ]).strip().upper()
            conn = get_connection()
            try:
                exists = conn.execute(
                    "SELECT code FROM airports WHERE code=%s", (code,)
                ).fetchone()
            finally:
                conn.close()
            if exists is None:
                errors.append(f"{champ} : code IATA '{code}' introuvable dans la base de donnees")

    return errors


@router.post("/import", response_model=FlightOut)
async def import_flight(_=Depends(get_current_user), file: UploadFile = File(...)):
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Format de fichier non supporte (attendu .xlsx)")

    content = await file.read()
    try:
        data = _read_excel_bytes(content)
    except Exception as e:
        raise HTTPException(400, f"Impossible de lire le fichier Excel : {e}")

    errors = _validate_excel(data)
    if errors:
        raise HTTPException(400, "; ".join(errors))

    body = FlightCreate(
        flight_number=str(data.get("Numero de vol", "IMPORT")),
        departure_airport=str(data.get("Aeroport depart", "")).upper(),
        arrival_airport=str(data.get("Aeroport arrivee", "")).upper(),
        distance_km=float(str(data.get("Distance (km)", 0)).replace(",", ".")),
        duration_hours=float(str(data.get("Duree (heures)", 0)).replace(",", ".")),
        aircraft_id=int(float(str(data.get("ID Avion", 1)))),
        passengers=int(float(str(data.get("Nombre de passagers", 0)).replace(",", "."))),
        fuel_price_per_liter=float(str(data.get("Prix carburant (MAD/L)", 0)).replace(",", ".")),
        ticket_price_avg=float(str(data.get("Prix billet moyen (MAD)", 0)).replace(",", ".")),
        catering_cost_per_pax=float(str(data.get("Catering par passager (MAD)", 25)).replace(",", ".")),
        handling_cost=float(str(data.get("Handling (MAD)", 500)).replace(",", ".")),
        taxes_airport=float(str(data.get("Taxes aeroportuaires (MAD)", 0)).replace(",", ".")),
        flight_date=str(data.get("Date du vol", "")) or None,
    )

    conn = get_connection()
    try:
        cur = conn.execute(
            f"INSERT INTO flights ({_COLUMNS}) VALUES ({_PLACEHOLDERS}) RETURNING id",
            _body_values(body),
        )
        new_id = cur.fetchone()["id"]
        conn.commit()
        row = _flight_row(conn, new_id)
    finally:
        conn.close()
    return _to_dict(row)
